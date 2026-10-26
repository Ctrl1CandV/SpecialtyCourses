#!/usr/bin/env python
"""Read-only structural QA for the C2-B optimized PPTX build.

This checker does not open Office/WPS, mutate files, render slides, or confirm
visual quality. WPS playback and visual review remain a separate manual gate.
"""
import argparse
import json
import os
import re
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')
import yaml
from pptx import Presentation

EMU_PER_INCH = 914400
NS = {
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}
STYLE_PATTERNS = [
    re.compile(r'立下.{0,4}(?:死|硬)?规矩'),
    re.compile(r'约定\s*[123一二三]'),
    re.compile(r'三条规矩'),
]


def lesson_path(raw_path):
    project_root = Path(__file__).resolve().parents[2]
    optimized_root = (project_root / '课程' / '课件' / 'C2' / 'C2-B-pro').resolve()
    lesson = Path(raw_path).resolve()
    if lesson.parent != optimized_root or not re.fullmatch(r'C2-B-(0[1-9]|10)', lesson.name):
        raise ValueError(f'只允许检查优化目录中的 C2-B-01 至 C2-B-10：{lesson}')
    return lesson


def duration_minutes(raw):
    match = re.match(r'^\s*(\d+)\s*(?:min|分钟)', str(raw), re.I)
    if not match:
        raise ValueError(f'无法读取课时：{raw!r}')
    return int(match.group(1))


def main():
    parser = argparse.ArgumentParser(description='只读检查优化版 C2-B 课件结构；不替代 WPS 实机验收。')
    parser.add_argument('lesson_dir')
    parser.add_argument('--soft', action='store_true', help='将措辞提示降为警告；结构错误仍返回失败。')
    args = parser.parse_args()

    lesson = lesson_path(args.lesson_dir)
    lesson_id = lesson.name
    errors, warnings = [], []
    deck_path = lesson / 'deck.yaml'
    script_path = lesson / f'{lesson_id}-讲稿.md'
    pptx_path = lesson / f'{lesson_id}.pptx'
    manifest_path = lesson / 'manifest.json'
    required = [deck_path, script_path, pptx_path, manifest_path]
    missing = [str(p.name) for p in required if not p.is_file()]
    if missing:
        print('缺少文件：' + ', '.join(missing))
        return 1

    deck = yaml.safe_load(deck_path.read_text(encoding='utf-8'))
    slides_spec = deck.get('slides') or []
    if not slides_spec:
        errors.append('deck.yaml 没有 slides')
    pages = {}
    pages_dir = lesson / 'pages'
    if pages_dir.exists():
        for page_file in sorted(pages_dir.glob('p*.yaml')):
            spec = yaml.safe_load(page_file.read_text(encoding='utf-8')) or {}
            page_no = spec.get('page')
            if not isinstance(page_no, int) or page_no in pages:
                errors.append(f'{page_file.name}: page 缺失、非整数或重复')
                continue
            pages[page_no] = spec

    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    manifest_slides = manifest.get('slides') or []
    prs = Presentation(str(pptx_path))
    n_slides = len(slides_spec)
    if len(prs.slides) != n_slides:
        errors.append(f'PPTX 页数 {len(prs.slides)} != deck.yaml 页数 {n_slides}')
    if len(manifest_slides) != n_slides:
        errors.append(f'manifest 页数 {len(manifest_slides)} != deck.yaml 页数 {n_slides}')
    if prs.slide_width / prs.slide_height < 1.77 or prs.slide_width / prs.slide_height > 1.78:
        errors.append(f'幻灯片比例不是 16:9：{prs.slide_width / prs.slide_height:.4f}')

    try:
        safe_bottom = float(deck.get('meta', {}).get('contentBottom', 6.42))
    except (TypeError, ValueError):
        safe_bottom = 6.42
        errors.append('meta.contentBottom 必须是英寸数')
    slide_w = prs.slide_width / EMU_PER_INCH
    slide_h = prs.slide_height / EMU_PER_INCH
    known_pages = set()

    for idx, slide in enumerate(prs.slides, 1):
        if idx > len(manifest_slides):
            continue
        man = manifest_slides[idx - 1]
        if man.get('page') != idx:
            errors.append(f'manifest 第 {idx} 页 page 字段错误：{man.get("page")}')
        entries = man.get('entries') or []
        shapes = list(slide.shapes)
        if len(shapes) != len(entries):
            errors.append(f'P{idx} PPTX 对象数 {len(shapes)} != manifest 条目数 {len(entries)}')
        by_name = {shape.name: shape for shape in shapes}
        if len(by_name) != len(shapes):
            errors.append(f'P{idx} 有重复形状名称')
        ids_by_key = {}
        page_no_entries = {i for i, entry in enumerate(entries) if (entry.get('name') or '').endswith(':pageno')}
        for shape_index, (shape, entry) in enumerate(zip(shapes, entries)):
            key = entry.get('id')
            if key:
                expected_name = key
                if shape.name != expected_name:
                    errors.append(f'P{idx} 形状应命名为 {expected_name}，实际为 {shape.name}')
                ids_by_key[key.split(':', 1)[-1]] = str(shape.shape_id)
            left = shape.left / EMU_PER_INCH
            top = shape.top / EMU_PER_INCH
            width = shape.width / EMU_PER_INCH
            height = shape.height / EMU_PER_INCH
            right, bottom = left + width, top + height
            if left < -0.02 or top < -0.02 or right > slide_w + 0.02 or bottom > slide_h + 0.02:
                errors.append(f'P{idx} 形状越出画布：({left:.2f},{top:.2f},{width:.2f},{height:.2f}) {shape.name}')
            elif shape_index not in page_no_entries and bottom > safe_bottom + 0.02:
                errors.append(f'P{idx} 内容侵入页脚安全区：底边 {bottom:.2f}in > {safe_bottom:.2f}in，{shape.name}')

        for rel in slide.part.rels.values():
            if 'image' in rel.reltype and not rel.is_external and not rel.target_part.blob:
                errors.append(f'P{idx} 图片关系 {rel.rId} 没有媒体内容')

        spec = pages.get(idx)
        if spec is None:
            continue
        known_pages.add(idx)
        steps = spec.get('steps') or []
        timing = slide._element.find('{%s}timing' % NS['p'])
        if steps and timing is None:
            errors.append(f'P{idx} 有动画步骤但 PPTX 缺少 timing')
            continue
        if not steps and timing is not None:
            errors.append(f'P{idx} 没有动画步骤却存在 timing')
            continue
        if not steps:
            continue

        seqs = timing.findall('.//p:seq', NS)
        if len(seqs) != 1:
            errors.append(f'P{idx} timing 中 seq 数量为 {len(seqs)}，应为 1')
            continue
        main_ctn = seqs[0].find('p:cTn', NS)
        group_list = main_ctn.find('p:childTnLst', NS) if main_ctn is not None else None
        groups = [child for child in (list(group_list) if group_list is not None else []) if child.tag == '{%s}par' % NS['p']]
        if len(groups) != len(steps):
            errors.append(f'P{idx} 点击组数 {len(groups)} != steps 数 {len(steps)}')

        animated_names = set()
        for element in (spec.get('anim') or []) + (spec.get('native') or []):
            if element.get('id'):
                animated_names.add(element['id'])
        for step_no, (step, group) in enumerate(zip(steps, groups), 1):
            expected_ids = set()
            for action in ('appear', 'hide'):
                for element_id in step.get(action, []):
                    key = f'p{idx:02d}:{element_id}'
                    if element_id not in animated_names:
                        errors.append(f'P{idx} steps[{step_no}].{action} 不能动画固定层或未定义元素：{key}')
                    elif key not in by_name:
                        errors.append(f'P{idx} steps[{step_no}].{action} 缺少形状：{key}')
                    else:
                        expected_ids.add(ids_by_key.get(element_id))
            actual_ids = {target.get('spid') for target in group.findall('.//p:spTgt', NS)}
            if expected_ids != actual_ids:
                errors.append(f'P{idx} 第 {step_no} 击目标不匹配：预计 {sorted(expected_ids)}，实际 {sorted(actual_ids)}')
        shape_ids = {str(shape.shape_id) for shape in shapes}
        for target in timing.findall('.//p:spTgt', NS):
            if target.get('spid') not in shape_ids:
                errors.append(f'P{idx} 动画引用不存在的形状 id={target.get("spid")}')
        timing_ids = [node.get('id') for node in timing.findall('.//p:cTn', NS)]
        if len(timing_ids) != len(set(timing_ids)):
            errors.append(f'P{idx} timing 的 cTn id 重复')

    for page_no in sorted(set(pages) - known_pages):
        errors.append(f'pages 中第 {page_no} 页超出 deck 页数')

    script = script_path.read_text(encoding='utf-8')
    headings = list(re.finditer(r'^## P(\d+)\b([^\n]*)$', script, re.M))
    page_nos = [int(match.group(1)) for match in headings]
    if page_nos != list(range(1, n_slides + 1)):
        errors.append(f'讲稿页序异常：{page_nos} != 1…{n_slides}')
    minutes = 0
    page_bodies = {}
    all_sections = list(re.finditer(r'^##\s+', script, re.M))
    for match in headings:
        page_no = int(match.group(1))
        header = match.group(2)
        time_match = re.search(r'（约\s*(\d+)\s*分钟）', header)
        no_voice = '不口播' in header
        if time_match:
            minutes += int(time_match.group(1))
        elif not no_voice:
            warnings.append(f'讲稿 P{page_no} 没有页级时长；请确认是否为不口播页')
        following = next((section.start() for section in all_sections if section.start() > match.start()), len(script))
        page_bodies[page_no] = script[match.end():following]
    try:
        expected_minutes = duration_minutes(deck.get('meta', {}).get('duration'))
        if minutes != expected_minutes:
            errors.append(f'讲稿页级时长合计 {minutes}min != meta.duration {expected_minutes}min')
    except ValueError as exc:
        errors.append(str(exc))

    for page_no, spec in pages.items():
        count = len(spec.get('steps') or [])
        if count:
            if f'【PPT动画·{count}击】' not in page_bodies.get(page_no, ''):
                warnings.append(f'讲稿 P{page_no} 未标注【PPT动画·{count}击】')

    style_findings = []
    for file_path in (deck_path, script_path):
        content = file_path.read_text(encoding='utf-8')
        for pattern in STYLE_PATTERNS:
            for match in pattern.finditer(content):
                excerpt = content[max(0, match.start() - 10):match.end() + 10].replace('\n', ' ')
                style_findings.append(f'{file_path.name}: 疑似生硬表述「{excerpt}」')
    if args.soft:
        warnings.extend(style_findings)
    else:
        errors.extend(style_findings)

    print(f'==== {lesson_id}：{n_slides} 页，混合规格页 {sorted(pages)} ====')
    for warning in warnings:
        print('  [警告] ' + warning)
    for error in errors:
        print('  [失败] ' + error)
    if errors:
        print(f'结构检查失败：{len(errors)} 项，{len(warnings)} 项警告')
        return 1
    print('结构检查通过；WPS 实际打开、视觉检查和逐击播放仍需人工验收。')
    return 0


if __name__ == '__main__':
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f'检查失败：{exc}', file=sys.stderr)
        raise SystemExit(1)
